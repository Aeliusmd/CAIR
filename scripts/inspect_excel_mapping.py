"""Inspect Excel mapping Modification column."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

PATH = Path(__file__).resolve().parents[1] / "Required Fields mapping with DB 2-11-2021_Modified (2)_updated.xlsx"


def main() -> int:
    wb = openpyxl.load_workbook(PATH)
    ws = wb["Sheet1"]
    for row in range(2, ws.max_row + 1):
        seg = (ws.cell(row, 1).value or "").strip()
        if not seg:
            continue
        mod = (ws.cell(row, 9).value or "").strip()
        dev = (ws.cell(row, 7).value or "").strip()
        mapping = str(ws.cell(row, 6).value or "").strip()
        print(f"R{row}: {seg}")
        print(f"  Mapping: {mapping[:120]}")
        if dev:
            print(f"  Dev: {dev[:120]}")
        print(f"  Mod: {mod[:250] if mod else '(empty)'}")
        print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
