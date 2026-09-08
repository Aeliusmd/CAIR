"""Fill Modification column in the Excel mapping workbook to match current code."""
from __future__ import annotations

from pathlib import Path

import openpyxl

PATH = Path(__file__).resolve().parents[1] / (
    "Required Fields mapping with DB 2-11-2021_Modified (2)_updated.xlsx"
)
FALLBACK_PATH = Path(__file__).resolve().parents[1] / (
    "Required Fields mapping with DB 2-11-2021_Modified (2)_updated_current.xlsx"
)

# Reflects current cair_integration code + verified CAIR onboarding tests (Sep 2026).
MODIFICATIONS: dict[str, str] = {
    "MSH-4": (
        "IMPLEMENTED: vendor/integration org SF-013259 via .env SENDING_FACILITY_ID. "
        "Used in MSH-4, PID-3.4, ORC placer/filler, SOAP facilityID. "
        "Different from clinic site org in MSH-22."
    ),
    "MSH-6": "IMPLEMENTED: CAIR2 via .env RECEIVING_FACILITY (not CAIRLO).",
    "MSH-15": "IMPLEMENTED: send AL (not ER). CAIR sample uses AL|AL.",
    "MSH-16": "IMPLEMENTED: send AL. CAIR sample uses AL|AL.",
    "MSH-22": (
        "IMPLEMENTED: clinic site org from .env PROVIDER_ORG_ID=SF-012218 "
        "(tested OK on CAIR stage). Must match RXA-11.4. "
        "Do NOT use vendor SF-013259 — CAIR rejects vendor in MSH-22. "
        "Sample CA0054321 also rejected. Longer term: Locations.CairOrgCode from DB."
    ),
    "PID-3-1.4": (
        "IMPLEMENTED: assigning authority = SENDING_FACILITY_ID (SF-013259). "
        "Not the EHR app name."
    ),
    "PID-3-1.5": "IMPLEMENTED: MR (Medical Record Number), not PI.",
    "PID-8": "IMPLEMENTED: GenderId map 1=M, 2=F, 3/Other=U in worker.",
    "PID-10.1": (
        "IMPLEMENTED: Patients.RaceId -> DataGroups.Description -> CDCREC "
        "(cdc_codes.py). Example White=2106-3^^CDCREC (RaceId 1082). "
        "CAIR warns if empty — fill RaceId in DB."
    ),
    "PID-11": (
        "IMPLEMENTED: Patients.Address1/City/State/ZipCode when all set. "
        "Sanitize skips placeholders (ADDRESS 1*). CAIR rejects non-street values "
        "(e.g. GAMPAHA). Use a real street line or leave blank."
    ),
    "PID-13": (
        "IMPLEMENTED: Patients.HomePhone, CellPhone, Email in PID-13 "
        "(PRN/PH, PRN/CP, NET/Internet). Required by CAIR email — do not leave empty."
    ),
    "PID-22-1": (
        "IMPLEMENTED: Patients.EthnicityId -> DataGroups.Description -> CDCREC. "
        "Example Not Hispanic or Latino=2186-5^^CDCREC (EthnicityId 1084). "
        "CAIR warns if empty — fill EthnicityId in DB."
    ),
    "PD1-12": (
        "IMPLEMENTED default N (shareable) per CAIR sample. No DB consent column yet. "
        "Y = lock from other providers."
    ),
    "PD1-13": (
        "IMPLEMENTED: EHRVaccines.VaccineDate (or visit date) when PD1-12 sent."
    ),
    "ORC-12": (
        "IMPLEMENTED: doctor on the visit (Providers via CheckInsHeader.ProviderId). "
        "XCN: NPI^Last^First^^^^^^NPPES&2.16.840.1.113883.4.6&ISO^^^^NPI^^^^^^^^MD "
        "(.9 assigning authority; .21 = professional suffix). "
        "HOW WE TAKE DEGREE AND TITLE (same for RXA-10.21): "
        "read Providers.Degree and Providers.Title; send ONE value only — "
        "(1) if Degree has a value, use Degree; "
        "(2) if Degree empty, use Title only when Title is a short credential "
        "(MD, DO, NP, RN, PA) — not a long job title like Physical Therapist; "
        "(3) else leave .21 blank. "
        "Example submission 3 (brian noor): Degree empty, Title=MD → ...^^^^^^^^MD. "
        "ACK MSA|AA — ORC-12.21 warning cleared."
    ),
    "RXA-10": (
        "IMPLEMENTED: same visit doctor as ORC-12 (Providers). "
        "RXA-10.21 uses the SAME Degree/Title rule as ORC-12.21: "
        "prefer Degree; if Degree empty use short Title (MD/NP/RN…); else blank. "
        "Example submission 3: Degree empty, Title=MD → ...^^^^^^^^MD. "
        "ACK MSA|AA — RXA-10.21 warning cleared."
    ),
    "RXA-3": (
        "IMPLEMENTED: EHRVaccines.VaccineDate + VaccineTime when set; "
        "else CheckInsHeader.CheckInDate + CheckInTime."
    ),
    "RXA-5": (
        "IMPLEMENTED: ServiceCodes.NDCNumber (via EHRVaccines.ServiceCodeId), "
        "dashed 11-digit form (e.g. 58160-0842-52^^NDC). "
        "CVX: ServiceCodes.CvxCode not in DB yet. CAIR accepts CVX or NDC, not both."
    ),
    "RXA-6": (
        "IMPLEMENTED: EHRVaccines.Dosage. If unknown send 999 with blank RXA-7. "
        "CAIR may still return informational RXA-6 warning for some values."
    ),
    "RXA-11.4": (
        "IMPLEMENTED: same as MSH-22 — .env PROVIDER_ORG_ID=SF-012218 "
        "(^^^SF-012218). Required when RXA-9.1=00. Not vendor SF-013259."
    ),
    "RXA-15": "IMPLEMENTED: EHRVaccines.LotNumber. Required when RXA-9=00.",
    "OBX-1": (
        "IMPLEMENTED: two OBX segments — OBX-1 VFC eligibility (64994-7 / V01); "
        "OBX-2 funding source (30963-3 / PHC70^Private^CDCPHINVS)."
    ),
    "OBX-5.1": "IMPLEMENTED: V01 (Not VFC eligible). Do not leave blank.",
    "OBX-11": "IMPLEMENTED: F (Final).",
}


def main() -> None:
    wb = openpyxl.load_workbook(PATH)
    ws = wb["Sheet1"]
    updated = 0
    for row in range(2, ws.max_row + 1):
        seg = (ws.cell(row, 1).value or "").strip()
        if seg not in MODIFICATIONS:
            continue
        ws.cell(row, 9).value = MODIFICATIONS[seg]
        updated += 1
    try:
        wb.save(PATH)
        print(f"Updated {updated} Modification cells -> {PATH}")
    except PermissionError:
        wb.save(FALLBACK_PATH)
        print(
            f"Updated {updated} Modification cells -> {FALLBACK_PATH}\n"
            f"(Original file locked — close Excel and replace/rename to the _updated.xlsx name)"
        )


if __name__ == "__main__":
    main()
